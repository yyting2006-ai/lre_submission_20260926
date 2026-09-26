#!/usr/bin/env python3
"""Reconcile the GF0025-2021 Appendix A inventory into 572 authority records.

The legacy workbook has 573 data rows because it includes one unnumbered
cross-reference note (legacy ``no=123``).  The official Appendix A gives that
note no grammar-point identifier.  This script excludes the note, expands the
workbook's merged taxonomy cells, assigns the official level-local identifiers,
and emits a versioned catalog plus a non-executable legacy-draft manifest.

No legacy regex is promoted to runtime by this script.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parents[1]
RESTRICTED_SOURCES = PROJECT_DIR / "restricted_sources"
DEFAULT_WORKBOOK = RESTRICTED_SOURCES / "语法_正则表达式v4.xlsx"
DEFAULT_OFFICIAL_PDF = RESTRICTED_SOURCES / "GF0025-2021.pdf"
DEFAULT_OCR_DIR = RESTRICTED_SOURCES / "gf0025_ocr"
DEFAULT_OUT_DIR = PROJECT_DIR / "kg"
SHEET_NAME = "语法_正则表达式 "
EXCLUDED_LEGACY_NO = 123

EXPECTED_LEVEL_COUNTS = {
    "1": 48,
    "2": 81,
    "3": 81,
    "4": 76,
    "5": 71,
    "6": 67,
    "7-9": 148,
}
OFFICIAL_LEVEL_PAGE_RANGES = {
    "1": (176, 184),
    "2": (185, 197),
    "3": (197, 211),
    "4": (211, 222),
    "5": (222, 233),
    "6": (233, 242),
    "7-9": (242, 259),
}
LEVEL_LABELS = {
    "一级": ("1", "一", "一级"),
    "二级": ("2", "二", "二级"),
    "三级": ("3", "三", "三级"),
    "四级": ("4", "四", "四级"),
    "五级": ("5", "五", "五级"),
    "六级": ("6", "六", "六级"),
    "高等": ("7-9", "七—九", "七—九级"),
    "七—九级": ("7-9", "七—九", "七—九级"),
}
CATALOG_FIELDS = [
    "standard_point_id",
    "official_label",
    "standard_code",
    "level",
    "level_label",
    "level_ordinal",
    "category",
    "subcategory",
    "section_item",
    "grammar_point",
    "canonical_label",
    "standard_text",
    "source_id",
    "source_label",
    "source_file_sha256",
    "source_pdf_page",
    "page_anchor",
    "page_anchor_status",
    "source_status",
    "authority_status",
    "legacy_workbook_no",
    "legacy_workbook_row",
    "legacy_regex",
    "legacy_note",
    "legacy_regex_status",
    "runtime_status",
    "evaluation_status",
]
RELEASE_FIELDS = [
    "standard_point_id",
    "official_label",
    "standard_code",
    "level",
    "level_label",
    "level_ordinal",
    "category",
    "subcategory",
    "section_item",
    "source_id",
    "source_file_sha256",
    "source_pdf_page",
    "page_anchor",
    "page_anchor_status",
    "authority_status",
    "legacy_workbook_no",
    "legacy_regex_status",
    "runtime_status",
    "evaluation_status",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--official-pdf", type=Path, default=DEFAULT_OFFICIAL_PDF)
    parser.add_argument("--ocr-dir", type=Path, default=DEFAULT_OCR_DIR)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT_DIR)
    return parser.parse_args()


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"[ \t]+", " ", text).strip()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def merged_value_map(ws: Any) -> dict[tuple[int, int], Any]:
    values: dict[tuple[int, int], Any] = {}
    for merged_range in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        top_left = ws.cell(min_row, min_col).value
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                values[(row, col)] = top_left
    return values


def workbook_rows(path: Path) -> tuple[list[dict[str, str]], list[str]]:
    wb = load_workbook(path, read_only=False, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Missing sheet {SHEET_NAME!r}; found {wb.sheetnames}")
    ws = wb[SHEET_NAME]
    merged = merged_value_map(ws)
    headers = [clean(ws.cell(1, col).value) for col in range(1, ws.max_column + 1)]
    # The workbook has one intentionally blank spacer column.  Give it a stable
    # name so dict construction remains lossless.
    headers = [header or f"_blank_col_{index}" for index, header in enumerate(headers, start=1)]
    rows: list[dict[str, str]] = []
    for row_number in range(2, ws.max_row + 1):
        payload: dict[str, str] = {"_workbook_row": str(row_number)}
        nonempty = False
        for col, header in enumerate(headers, start=1):
            value = merged.get((row_number, col), ws.cell(row_number, col).value)
            payload[header] = clean(value)
            nonempty = nonempty or bool(payload[header])
        if nonempty:
            rows.append(payload)
    return rows, headers


def canonical_label(standard_text: str) -> str:
    first_line = standard_text.split("\n", 1)[0].strip()
    if "：" in first_line:
        head, tail = first_line.split("：", 1)
        if head and len(head) <= 32:
            return head.strip()
        return f"{head.strip()}：{tail.strip()}".strip("：")
    return first_line


def compact_ocr(text: str) -> str:
    compact = re.sub(r"\s+", "", text)
    return (
        compact.replace("－", "—")
        .replace("–", "—")
        .replace("七一九", "七—九")
        .replace("七-九", "七—九")
    )


def ocr_page_anchors(ocr_dir: Path) -> tuple[dict[tuple[str, int], int], dict[str, Any]]:
    anchors: dict[tuple[str, int], int] = {}
    duplicates: dict[str, list[int]] = defaultdict(list)
    if not ocr_dir.exists():
        return anchors, {"ocr_directory": str(ocr_dir), "available": False, "matched_ids": 0}

    # Windows OCR often drops the closing bracket while retaining the opening
    # bracket, level, and ordinal. Match that stable prefix, then choose the
    # earliest occurrence that preserves the official ordinal sequence within
    # each level. This rejects later backward references and most forward
    # cross-references without inventing a page.
    bracket_pattern = re.compile(
        r"[〔【［\[](?P<level>七—九|[一二三四五六])(?P<ordinal>\d{1,3})"
    )
    char_to_level = {"一": "1", "二": "2", "三": "3", "四": "4", "五": "5", "六": "6", "七—九": "7-9"}
    candidates: dict[tuple[str, int], list[tuple[int, int]]] = defaultdict(list)
    for text_path in sorted(ocr_dir.glob("page-*.txt")):
        match = re.search(r"page-(\d+)$", text_path.stem)
        if not match:
            continue
        page = int(match.group(1))
        text = compact_ocr(text_path.read_text(encoding="utf-8"))
        for point_match in bracket_pattern.finditer(text):
            level = char_to_level[point_match.group("level")]
            ordinal = int(point_match.group("ordinal"))
            if ordinal < 1 or ordinal > EXPECTED_LEVEL_COUNTS[level]:
                continue
            start_page, end_page = OFFICIAL_LEVEL_PAGE_RANGES[level]
            if start_page <= page <= end_page:
                candidates[(level, ordinal)].append((page, point_match.start()))

    for level, expected_count in EXPECTED_LEVEL_COUNTS.items():
        previous_position = (OFFICIAL_LEVEL_PAGE_RANGES[level][0], -1)
        for ordinal in range(1, expected_count + 1):
            options = sorted(set(candidates.get((level, ordinal), [])))
            eligible = [position for position in options if position > previous_position]
            if not eligible:
                continue
            chosen = eligible[0]
            anchors[(level, ordinal)] = chosen[0]
            previous_position = chosen
            if len(options) > 1:
                duplicates[f"{level}:{ordinal}"] = [page for page, _ in options]
    return anchors, {
        "ocr_directory": str(ocr_dir.resolve()),
        "available": True,
        "matched_ids": len(anchors),
        "candidate_token_matches": sum(len(value) for value in candidates.values()),
        "duplicate_page_matches": {key: sorted(set(value)) for key, value in duplicates.items()},
    }


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    args = parse_args()
    workbook = args.workbook.resolve()
    official_pdf = args.official_pdf.resolve()
    out_dir = args.out_dir.resolve()
    if not workbook.exists():
        raise FileNotFoundError(workbook)
    if not official_pdf.exists():
        raise FileNotFoundError(official_pdf)

    raw_rows, headers = workbook_rows(workbook)
    if len(raw_rows) != 573:
        raise ValueError(f"Expected 573 legacy rows before reconciliation, found {len(raw_rows)}")

    excluded = [row for row in raw_rows if int(row["no"]) == EXCLUDED_LEGACY_NO]
    if len(excluded) != 1:
        raise ValueError(f"Expected exactly one legacy no={EXCLUDED_LEGACY_NO} row")
    excluded_text = excluded[0].get("grammar_content", "")
    if not excluded_text.startswith("※") or "是……的" not in excluded_text:
        raise ValueError(f"Excluded row does not match the verified cross-reference note: {excluded_text}")
    rows = [row for row in raw_rows if int(row["no"]) != EXCLUDED_LEGACY_NO]
    if len(rows) != 572:
        raise AssertionError(len(rows))

    source_hash = sha256_file(official_pdf)
    anchors, ocr_report = ocr_page_anchors(args.ocr_dir.resolve())
    ordinals: Counter[str] = Counter()
    catalog: list[dict[str, Any]] = []
    for row in rows:
        raw_level = row.get("level", "")
        if raw_level not in LEVEL_LABELS:
            raise ValueError(f"Unknown level {raw_level!r} at workbook row {row['_workbook_row']}")
        level, official_level, level_label = LEVEL_LABELS[raw_level]
        ordinals[level] += 1
        ordinal = ordinals[level]
        machine_id = f"GF0025-2021-{level.replace('-', '_')}-{ordinal:03d}"
        official_label = f"【{official_level}{ordinal:02d}】"
        standard_text = row.get("grammar_content", "")
        if not standard_text:
            raise ValueError(f"Empty grammar_content at workbook row {row['_workbook_row']}")
        legacy_regex = row.get("正则表达式", "")
        pdf_page = anchors.get((level, ordinal))
        anchor_status = "ocr_detected" if pdf_page else "appendix_span_verified"
        catalog.append(
            {
                "standard_point_id": machine_id,
                "official_label": official_label,
                "standard_code": "GF 0025—2021 Appendix A",
                "level": level,
                "level_label": level_label,
                "level_ordinal": ordinal,
                "category": row.get("grammar_item", ""),
                # Four top-level categories have no named lower subdivision in
                # the legacy sheet.  They share an explicit sentinel instead
                # of being emitted as empty taxonomy values.  This preserves
                # the Standard's 33 distinct subitem labels.
                "subcategory": row.get("category", "") or "未细分",
                "section_item": row.get("ximu", ""),
                "grammar_point": standard_text,
                "canonical_label": canonical_label(standard_text),
                "standard_text": standard_text,
                "source_id": "REF-GF0025-2021",
                "source_label": "GF 0025—2021 Chinese Proficiency Grading Standards: Appendix A",
                "source_file_sha256": source_hash,
                "source_pdf_page": str(pdf_page or ""),
                "page_anchor": f"GF0025-2021 PDF p.{pdf_page}" if pdf_page else "GF0025-2021 PDF pp.176–259",
                "page_anchor_status": anchor_status,
                "source_status": "official_numbered_entry",
                "authority_status": "verified_inventory",
                "legacy_workbook_no": row.get("no", ""),
                "legacy_workbook_row": row.get("_workbook_row", ""),
                "legacy_regex": legacy_regex,
                "legacy_note": row.get("备注", ""),
                "legacy_regex_status": "draft_unreviewed" if legacy_regex else "not_available",
                "runtime_status": "not_promoted",
                "evaluation_status": "outside_current_37_label_claim_unless_crosswalked",
            }
        )

    if dict(ordinals) != EXPECTED_LEVEL_COUNTS:
        raise ValueError(f"Level distribution mismatch: {dict(ordinals)}")
    ids = [row["standard_point_id"] for row in catalog]
    if len(ids) != len(set(ids)):
        raise ValueError("Duplicate standard_point_id")
    official_labels = [row["official_label"] for row in catalog]
    if len(official_labels) != len(set(official_labels)):
        raise ValueError("Duplicate official_label")
    if len({row["category"] for row in catalog}) != 12:
        raise ValueError("Expanded taxonomy does not contain exactly 12 primary categories")
    if len({row["subcategory"] for row in catalog}) != 33:
        raise ValueError("Expanded taxonomy does not contain exactly 33 subitem labels")

    out_dir.mkdir(parents=True, exist_ok=True)
    catalog_path = out_dir / "standard_grammar_points_gf0025_v2.csv"
    release_catalog_path = out_dir / "standard_inventory_release_v2.csv"
    excluded_path = out_dir / "standard_inventory_excluded_notes_v2.csv"
    drafts_path = out_dir / "diagnostic_rule_drafts_gf0025_v2.jsonl"
    report_path = out_dir / "full_inventory_reconciliation_v2.json"
    write_csv(catalog_path, catalog, CATALOG_FIELDS)
    write_csv(release_catalog_path, catalog, RELEASE_FIELDS)
    write_csv(
        excluded_path,
        [
            {
                "legacy_workbook_no": row["no"],
                "legacy_workbook_row": row["_workbook_row"],
                "text": row.get("grammar_content", ""),
                "exclusion_reason": "unnumbered_cross_reference_note_in_official_appendix",
            }
            for row in excluded
        ],
        ["legacy_workbook_no", "legacy_workbook_row", "text", "exclusion_reason"],
    )

    with drafts_path.open("w", encoding="utf-8") as handle:
        for row in catalog:
            payload = {
                "standard_point_id": row["standard_point_id"],
                "official_label": row["official_label"],
                "canonical_label": row["canonical_label"],
                "legacy_surface_pattern": row["legacy_regex"],
                "draft_status": row["legacy_regex_status"],
                "runtime_eligible": False,
                "promotion_requirements": [
                    "source_page_audit",
                    "typed_slot_compilation",
                    "linguist_review",
                    "positive_behavioral_tests",
                    "hard_negative_behavioral_tests",
                ],
            }
            handle.write(json.dumps(payload, ensure_ascii=False) + "\n")

    regex_counts = Counter(row["legacy_regex_status"] for row in catalog)
    anchor_counts = Counter(row["page_anchor_status"] for row in catalog)
    report = {
        "version": "v2.0",
        "standard": "GF 0025—2021",
        "workbook": str(workbook),
        "workbook_headers": headers,
        "official_pdf": str(official_pdf),
        "official_pdf_sha256": source_hash,
        "legacy_rows": len(raw_rows),
        "excluded_rows": 1,
        "excluded_legacy_no": EXCLUDED_LEGACY_NO,
        "excluded_reason": "unnumbered cross-reference note between official 二74 and 二75",
        "catalog_rows": len(catalog),
        "unique_standard_point_ids": len(set(ids)),
        "level_distribution": dict(ordinals),
        "expected_level_distribution": EXPECTED_LEVEL_COUNTS,
        "primary_categories": len({row["category"] for row in catalog}),
        "subitem_labels": len({row["subcategory"] for row in catalog}),
        "category_subitem_memberships": len(
            {(row["category"], row["subcategory"]) for row in catalog}
        ),
        "legacy_regex_status_counts": dict(regex_counts),
        "page_anchor_status_counts": dict(anchor_counts),
        "ocr": ocr_report,
        "scope": {
            "authority_inventory_complete": True,
            "legacy_regexes_promoted_to_runtime": 0,
            "runtime_or_evaluation_claim": False,
        },
        "outputs": {
            "catalog": str(catalog_path),
            "release_safe_catalog": str(release_catalog_path),
            "excluded_note": str(excluded_path),
            "diagnostic_drafts": str(drafts_path),
        },
        "integrity_pass": True,
    }
    write_json(report_path, report)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
